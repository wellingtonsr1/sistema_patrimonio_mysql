"""
Testes do módulo de Inventário Patrimonial Completo e Comprobatório.

Cobre (prompt §19): criação, escopo/snapshot, conferência, divergências,
bem não previsto, encerramento, RBAC, auditoria e regressão do cadastro
(o inventário NUNCA altera o patrimônio).
"""

import pytest
from datetime import datetime

from app.models.enums import AssetCategory, AssetCondition, AssetStatus, InventarioStatus, InventarioItemStatus
from app.schemas.asset import AssetCreate
from app.services.asset_service import AssetService
from app.services.location_service import LocationService
from app.schemas.location import LocationCreate
from app.services.inventario_service import InventarioService
from app.services.report_service import ReportService
from app.services.permission_service import assign_role, get_role_by_name, ensure_default_roles
from app.services.auth_service import create_user
from app.models.audit_log import AuditLog


def _make_location(db, name="Matriz - TI - Sala dos Servidores", department="Tecnologia da Informação"):
    return LocationService.create(db, LocationCreate(
        name=name, branch="Matriz", department=department,
    ))


def _make_asset(db, tag, location=None, name="Notebook Dell"):
    return AssetService.create(db, AssetCreate(
        tag=tag,
        name=name,
        category=AssetCategory.NOTEBOOK,
        purchase_value=3500.0,
        condition=AssetCondition.GOOD,
        initial_location_id=location.id if location else None,
    ))


# ============================================================================
# Serviço — criação e escopo
# ============================================================================

def test_create_inventario_generates_expected_items_with_snapshot(db_session):
    loc = _make_location(db_session)
    a1 = _make_asset(db_session, "INV-TST-0001", location=loc)
    _make_asset(db_session, "INV-TST-0002", location=loc)
    _make_asset(db_session, "INV-TST-0003", location=None)  # fora do escopo (sem local)

    inv = InventarioService.create_inventario(
        db_session, name="Inventário TI 2026", location_id=loc.id,
        created_by_name="admin",
    )

    assert inv.code.startswith("INV-")
    assert inv.status == InventarioStatus.PLANNED
    assert "Todo o acervo" not in inv.scope_filters and loc.name in inv.scope_filters
    tags = {i.asset.tag for i in inv.itens}
    assert tags == {"INV-TST-0001", "INV-TST-0002"}

    item1 = [i for i in inv.itens if i.asset.tag == "INV-TST-0001"][0]
    assert item1.status == InventarioItemStatus.PENDING
    assert item1.expected_location_name == loc.name
    assert item1.nao_previsto is False
    assert a1.location_id == loc.id  # cadastro intacto


def test_scope_snapshot_survives_later_asset_move(db_session):
    """O snapshot da expectativa não muda quando o bem é movimentado depois."""
    loc_a = _make_location(db_session, name="Sede A", department="TI")
    loc_b = _make_location(db_session, name="Sede B", department="RH")
    asset = _make_asset(db_session, "INV-TST-0010", location=loc_a)
    inv = InventarioService.create_inventario(db_session, name="Snapshot", location_id=loc_a.id)

    # Movimenta o bem para outro local (fluxo normal do sistema)
    asset.location_id = loc_b.id
    db_session.commit()

    item = inv.itens[0]
    assert item.expected_location_name == loc_a.name
    assert item.expected_location_id == loc_a.id


# ============================================================================
# Serviço — conferência
# ============================================================================

def _inv_with_one_item(db, tag="INV-TST-0020"):
    loc = _make_location(db, name="Loc Conferência", department="TI")
    asset = _make_asset(db, tag, location=loc)
    inv = InventarioService.create_inventario(db, name="Conf", location_id=loc.id)
    return inv, asset, loc


def test_record_check_found_starts_inventory(db_session):
    inv, asset, loc = _inv_with_one_item(db_session)
    item = inv.itens[0]

    assert inv.started_at is None
    status_before = asset.status
    result = InventarioService.record_check(
        db_session, item=item, result=InventarioItemStatus.FOUND,
        user_id=1, username="conferente",
    )

    assert result.status == InventarioItemStatus.FOUND
    assert result.checked_by_name == "conferente"
    assert result.checked_at is not None
    assert result.found_location_name == loc.name
    inv = InventarioService.get_by_id(db_session, inv.id)
    assert inv.started_at is not None
    assert inv.status == InventarioStatus.IN_PROGRESS
    # Cadastro do bem permanece intacto (regra §10)
    db_session.expire_all()
    assert asset.location_id == loc.id
    assert asset.status == status_before


def test_record_check_wrong_location_requires_different_location(db_session):
    inv, asset, loc = _inv_with_one_item(db_session, tag="INV-TST-0021")
    other_loc = _make_location(db_session, name="Outro Local", department="RH")
    item = inv.itens[0]

    # Mesmo local do esperado não pode ser registrado como divergência
    with pytest.raises(ValueError):
        InventarioService.record_check(
            db_session, item=item,
            result=InventarioItemStatus.FOUND_WRONG_LOCATION,
            found_location_id=loc.id,
        )
    # Sem local informado também não
    with pytest.raises(ValueError):
        InventarioService.record_check(
            db_session, item=item,
            result=InventarioItemStatus.FOUND_WRONG_LOCATION,
        )

    InventarioService.record_check(
        db_session, item=item,
        result=InventarioItemStatus.FOUND_WRONG_LOCATION,
        found_location_id=other_loc.id,
        observation="Bem encontrado na sala vizinha",
    )
    assert item.status == InventarioItemStatus.FOUND_WRONG_LOCATION
    assert item.found_location_name == other_loc.name
    assert item.observation == "Bem encontrado na sala vizinha"
    # O inventário NÃO altera o cadastro do bem (regra §10)
    assert asset.location_id == loc.id


def test_record_check_not_found(db_session):
    inv, asset, loc = _inv_with_one_item(db_session, tag="INV-TST-0022")
    item = inv.itens[0]
    InventarioService.record_check(
        db_session, item=item, result=InventarioItemStatus.NOT_FOUND,
        username="conferente",
    )
    assert item.status == InventarioItemStatus.NOT_FOUND
    assert item.found_location_id is None


def test_record_check_blocked_after_close(db_session):
    inv, asset, loc = _inv_with_one_item(db_session, tag="INV-TST-0023")
    item = inv.itens[0]
    InventarioService.record_check(db_session, item=item, result=InventarioItemStatus.FOUND)
    InventarioService.close_inventario(db_session, inventario=inv, closed_by_name="admin")

    with pytest.raises(ValueError):
        InventarioService.record_check(
            db_session, item=item, result=InventarioItemStatus.NOT_FOUND
        )


# ============================================================================
# Serviço — bem não previsto e encerramento
# ============================================================================

def test_register_unlisted_asset_and_duplicate_guard(db_session):
    inv, asset_in, loc = _inv_with_one_item(db_session, tag="INV-TST-0030")
    outsider = _make_asset(db_session, "INV-TST-0031", location=None)  # fora do escopo

    item, created = InventarioService.register_unlisted_asset(
        db_session, inventario=inv, asset_id=outsider.id,
        found_location_id=loc.id, observation="Bem esquecido no canto",
        username="conferente",
    )
    assert created is True
    assert item.nao_previsto is True
    assert item.found_location_name == loc.name

    # Segundo registro do mesmo bem não duplica
    item2, created2 = InventarioService.register_unlisted_asset(
        db_session, inventario=inv, asset_id=outsider.id,
    )
    assert created2 is False
    assert item2.id == item.id

    summary = InventarioService.summary(db_session, inv.id)
    assert summary["expected"] == 1        # esperados não mudam
    assert summary["unlisted"] == 1


def test_close_requires_all_expected_checked(db_session):
    loc = _make_location(db_session, name="Loc Encerrar", department="TI")
    _make_asset(db_session, "INV-TST-0040", location=loc)
    _make_asset(db_session, "INV-TST-0041", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Encerramento", location_id=loc.id)

    # Com pendentes, não encerra
    with pytest.raises(ValueError):
        InventarioService.close_inventario(db_session, inventario=inv)

    for item in inv.itens:
        InventarioService.record_check(db_session, item=item, result=InventarioItemStatus.FOUND)

    closed = InventarioService.close_inventario(
        db_session, inventario=inv, closed_by_name="admin", closure_notes="OK"
    )
    assert closed.status == InventarioStatus.CLOSED
    assert closed.closed_at is not None
    assert closed.closed_by_name == "admin"

    # Duplo encerramento não é permitido
    with pytest.raises(ValueError):
        InventarioService.close_inventario(db_session, inventario=inv)


def test_summary_counts(db_session):
    loc = _make_location(db_session, name="Loc Resumo", department="TI")
    _make_asset(db_session, "INV-TST-0050", location=loc)
    _make_asset(db_session, "INV-TST-0051", location=loc)
    _make_asset(db_session, "INV-TST-0052", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Resumo", location_id=loc.id)
    itens = {i.asset.tag: i for i in inv.itens}

    InventarioService.record_check(db_session, item=itens["INV-TST-0050"], result=InventarioItemStatus.FOUND)
    InventarioService.record_check(db_session, item=itens["INV-TST-0051"], result=InventarioItemStatus.NOT_FOUND)

    s = InventarioService.summary(db_session, inv.id)
    assert s["expected"] == 3
    assert s["found"] == 1
    assert s["not_found"] == 1
    assert s["pending"] == 1
    assert s["checked"] == 2


# ============================================================================
# Web — fluxos HTTP, RBAC e auditoria
# ============================================================================

def test_inventario_pages_require_login(unauth_client):
    resp = unauth_client.get("/inventarios", follow_redirects=False)
    assert resp.status_code == 303
    assert "/login" in resp.headers["Location"]


def test_full_inventory_flow_via_web(client, db_session):
    loc = _make_location(db_session, name="Loc Web", department="TI")
    asset = _make_asset(db_session, "INV-WEB-0001", location=loc)

    # Criar inventário (form)
    resp = client.post("/inventarios/new", data={"name": "Inventário Web", "location_id": str(loc.id)}, follow_redirects=False)
    assert resp.status_code == 303, resp.text
    detail_url = resp.headers["Location"]
    assert "/inventarios/" in detail_url
    inv_id = int(detail_url.rsplit("/", 1)[1])

    # Listagem e detalhe
    assert client.get("/inventarios").status_code == 200
    detail = client.get(detail_url)
    assert detail.status_code == 200
    assert "INV-WEB-0001" in detail.text

    # Localizar bem pelo tombamento (fluxo de campo)
    resp = client.post(f"/inventarios/{inv_id}/buscar", data={"tag": "INV-WEB-0001"}, follow_redirects=False)
    assert resp.status_code == 303

    # Item do bem
    inv = InventarioService.get_by_id(db_session, inv_id)
    item = [i for i in inv.itens if i.asset.tag == "INV-WEB-0001"][0]

    # Conferir via POST (rota web)
    resp = client.post(
        f"/inventarios/{inv_id}/conferir/{item.id}",
        data={"result": "ENCONTRADO", "observation": "Conferido no Web"},
        follow_redirects=False,
    )
    assert resp.status_code == 303
    db_session.expire_all()
    item = InventarioService.get_item(db_session, item.id)
    assert item.status == InventarioItemStatus.FOUND

    # Encerrar
    resp = client.post(f"/inventarios/{inv_id}/encerrar", data={"closure_notes": "Concluído"}, follow_redirects=False)
    assert resp.status_code == 303
    inv = InventarioService.get_by_id(db_session, inv_id)
    assert inv.status == InventarioStatus.CLOSED

    # Auditoria registrada (criação + conferência + encerramento)
    logs = db_session.query(AuditLog).filter(AuditLog.module == "Inventário").all()
    actions = {log.description for log in logs}
    assert any("criado" in d for d in actions)
    assert any("Conferência registrada" in d for d in actions)
    assert any("encerrado" in d for d in actions)


def test_web_close_blocked_with_pending(client, db_session):
    loc = _make_location(db_session, name="Loc Web Pend", department="TI")
    _make_asset(db_session, "INV-WEB-0002", location=loc)
    resp = client.post("/inventarios/new", data={"name": "Pendentes", "location_id": str(loc.id)}, follow_redirects=False)
    inv_id = int(resp.headers["Location"].rsplit("/", 1)[1])

    resp = client.post(f"/inventarios/{inv_id}/encerrar", data={}, follow_redirects=False)
    assert resp.status_code == 303
    assert "error=" in resp.headers["Location"]


def test_rbac_deny_and_grant_inventario(client, db_session):
    """Usuário sem permissão → 403; com perfil Patrimônio → acessa."""
    ensure_default_roles(db_session)
    plain = create_user(db_session, username="plainuser", password="teste@1234", full_name="Plain")
    client.post("/api/v1/auth/login", data={"username": "plainuser", "password": "teste@1234"})
    assert client.get("/inventarios").status_code == 403

    role = get_role_by_name(db_session, "Patrimônio")
    assign_role(db_session, plain, role)
    assert client.get("/inventarios").status_code == 200

    # Volta para o admin (não vaza sessão para outros testes)
    client.post("/api/v1/auth/login", data={"username": "testuser", "password": "teste@1234"})


def test_asset_detail_offers_inventory_conference(client, db_session):
    """QR landing: ficha do bem passa a oferecer conferência p/ inventários abertos."""
    loc = _make_location(db_session, name="Loc QR", department="TI")
    asset = _make_asset(db_session, "INV-QR-0001", location=loc)
    client.post(
        "/inventarios/new",
        data={"name": "QR Flow", "location_id": str(loc.id)},
        follow_redirects=False,
    )
    page = client.get(f"/assets/{asset.id}")
    assert page.status_code == 200
    assert "Inventário" in page.text
    assert "/inventarios/" in page.text

    # Página de conferência do bem
    from app.models.inventario import Inventario as Inv
    inv = db_session.query(Inv).first()
    resp = client.get(f"/inventarios/{inv.id}/conferir/{asset.id}")
    assert resp.status_code == 200
    assert "INV-QR-0001" in resp.text
    assert "Encontrado" in resp.text


# ============================================================================
# EXPORTAÇÃO DA ATA (CSV / PDF)
# ============================================================================

def test_export_inventario_csv_content(client, db_session):
    """A ata em CSV contém comprovação, snapshot esperado e resultado conferido."""
    loc = _make_location(db_session, name="Loc Export CSV", department="TI")
    other_loc = _make_location(db_session, name="Local Divergente", department="RH")
    a1 = _make_asset(db_session, "INV-EXP-0001", location=loc)
    _make_asset(db_session, "INV-EXP-0002", location=loc)
    outsider = _make_asset(db_session, "INV-EXP-0003", location=None)

    inv = InventarioService.create_inventario(db_session, name="Export CSV", location_id=loc.id)
    itens = {i.asset.tag: i for i in inv.itens}
    InventarioService.record_check(
        db_session, item=itens["INV-EXP-0001"], result=InventarioItemStatus.FOUND,
        user_id=1, username="conferente1",
    )
    InventarioService.record_check(
        db_session, item=itens["INV-EXP-0002"],
        result=InventarioItemStatus.FOUND_WRONG_LOCATION,
        found_location_id=other_loc.id, username="conferente2",
    )
    InventarioService.register_unlisted_asset(
        db_session, inventario=inv, asset_id=outsider.id,
        found_location_id=loc.id, username="conferente1",
    )

    csv_text = ReportService.generate_inventario_csv(db_session, inv)

    # Bloco de comprovação do cabeçalho
    assert "ATA DE INVENTÁRIO PATRIMONIAL" in csv_text
    assert inv.code in csv_text
    assert "Export CSV" in csv_text
    assert loc.name in csv_text  # escopo

    # Dados por item
    assert "INV-EXP-0001" in csv_text
    assert "Encontrado" in csv_text
    assert "Local Diferente" in csv_text
    assert "conferente1" in csv_text
    assert "conferente2" in csv_text
    assert "Bem não previsto" in csv_text
    assert "INV-EXP-0003" in csv_text

    # Snapshot do esperado preservado mesmo com divergência
    assert f"{loc.name}" in csv_text
    assert "Local Divergente" in csv_text


def test_export_inventario_pdf_bytes(client, db_session):
    """PDF gerado é um documento válido (assinatura %PDF) e usa a fonte de dados correta."""
    loc = _make_location(db_session, name="Loc Export PDF", department="TI")
    _make_asset(db_session, "INV-EXP-0011", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Export PDF", location_id=loc.id)
    InventarioService.record_check(
        db_session, item=inv.itens[0], result=InventarioItemStatus.FOUND,
        user_id=1, username="conferente",
    )

    pdf_bytes = ReportService.generate_inventario_pdf(db_session, inv)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 500  # documento com conteúdo, não página em brano


def test_export_endpoints_require_export_permission(db_session, unauth_client):
    """Usuário com inventario.visualizar (Consulta) mas sem relatorios.exportar → 403."""
    ensure_default_roles(db_session)
    from tests.test_rbac import _make_user, _login, PASSWORD
    _make_user(db_session, "consulta_exp", role_names=["Consulta"])
    _login(unauth_client, "consulta_exp", PASSWORD)

    loc = _make_location(db_session, name="Loc RBAC Exp", department="TI")
    _make_asset(db_session, "INV-EXP-0021", location=loc)
    inv = InventarioService.create_inventario(db_session, name="RBAC Exp", location_id=loc.id)

    assert unauth_client.get(f"/api/v1/reports/inventarios/{inv.id}/csv").status_code == 403
    assert unauth_client.get(f"/api/v1/reports/inventarios/{inv.id}/excel").status_code == 403
    assert unauth_client.get(f"/api/v1/reports/inventarios/{inv.id}/pdf").status_code == 403


def test_export_endpoints_404_unknown_inventario(client):
    assert client.get("/api/v1/reports/inventarios/999999/csv").status_code == 404
    assert client.get("/api/v1/reports/inventarios/999999/excel").status_code == 404
    assert client.get("/api/v1/reports/inventarios/999999/pdf").status_code == 404


def test_export_via_http_end_to_end(client, db_session):
    """Fluxo HTTP completo: admin (todas as permissões) baixa CSV e PDF."""
    loc = _make_location(db_session, name="Loc HTTP Exp", department="TI")
    asset = _make_asset(db_session, "INV-EXP-0031", location=loc)
    inv = InventarioService.create_inventario(db_session, name="HTTP Exp", location_id=loc.id)
    InventarioService.record_check(
        db_session, item=inv.itens[0], result=InventarioItemStatus.FOUND,
        user_id=1, username="admin",
    )

    csv_resp = client.get(f"/api/v1/reports/inventarios/{inv.id}/csv")
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-type"].startswith("text/csv")
    assert f"ata_{inv.code}.csv" in csv_resp.headers["content-disposition"]
    assert asset.tag in csv_resp.text
    assert "ATA DE INVENTÁRIO PATRIMONIAL" in csv_resp.text

    pdf_resp = client.get(f"/api/v1/reports/inventarios/{inv.id}/pdf")
    assert pdf_resp.status_code == 200
    assert pdf_resp.headers["content-type"] == "application/pdf"
    assert pdf_resp.content.startswith(b"%PDF-")


def test_export_inventario_excel_roundtrip(db_session):
    """XLSX gerado é uma planilha válida com comprovação, tabela e dados conferidos."""
    from openpyxl import load_workbook
    import io as _io

    loc = _make_location(db_session, name="Loc Export XLSX", department="TI")
    other_loc = _make_location(db_session, name="Local Divergente XLSX", department="RH")
    _make_asset(db_session, "INV-EXP-0041", location=loc)
    _make_asset(db_session, "INV-EXP-0042", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Export XLSX", location_id=loc.id)
    itens = {i.asset.tag: i for i in inv.itens}
    InventarioService.record_check(
        db_session, item=itens["INV-EXP-0041"], result=InventarioItemStatus.FOUND,
        user_id=1, username="conferente1",
    )
    InventarioService.record_check(
        db_session, item=itens["INV-EXP-0042"],
        result=InventarioItemStatus.FOUND_WRONG_LOCATION,
        found_location_id=other_loc.id, username="conferente2",
        observation="Encontrado no setor vizinho",
    )

    xlsx_bytes = ReportService.generate_inventario_excel(db_session, inv)
    assert isinstance(xlsx_bytes, bytes)
    # Assinatura ZIP (formato .xlsx é um pacite OOXML)
    assert xlsx_bytes[:2] == b"PK"

    wb = load_workbook(filename=_io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title.startswith("Ata INV-")

    # Bloco de comprovação
    assert ws.cell(row=1, column=1).value == "ATA DE INVENTÁRIO PATRIMONIAL"
    assert ws.cell(row=2, column=2).value == inv.code
    assert ws.cell(row=3, column=2).value == "Export XLSX"

    # Cabeçalho da tabela de itens
    all_values = {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}
    assert "Tombamento" in all_values
    assert "Resultado" in all_values
    assert "INV-EXP-0041" in all_values
    assert "INV-EXP-0042" in all_values
    assert "Encontrado" in all_values
    assert "Local Diferente" in all_values
    assert "conferente1" in all_values
    assert "conferente2" in all_values
    assert "Encontrado no setor vizinho" in all_values
    assert "Local Divergente XLSX" in all_values  # local encontrado (divergência)


def test_export_inventario_excel_via_http(client, db_session):
    """Endpoint HTTP: content-type e filename corretos para .xlsx."""
    loc = _make_location(db_session, name="Loc XLSX HTTP", department="TI")
    _make_asset(db_session, "INV-EXP-0051", location=loc)
    inv = InventarioService.create_inventario(db_session, name="XLSX HTTP", location_id=loc.id)

    resp = client.get(f"/api/v1/reports/inventarios/{inv.id}/excel")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"ata_{inv.code}.xlsx" in resp.headers["content-disposition"]
    assert resp.content[:2] == b"PK"


# ============================================================================
# Feature 034 — Snapshot de novos inventários sem colaborador responsável (US1)
# ============================================================================

def _assign_custodian(db, asset, name, email):
    """Cria colaborador e o atribui ao bem pelo fluxo oficial de alocação."""
    from app.schemas.custodian import CustodianCreate
    from app.services.custodian_service import CustodianService
    from app.services.movement_service import MovementService
    from app.schemas.movement import MovementCreate
    from app.models.enums import MovementType

    cust = CustodianService.create(db, CustodianCreate(
        name=name, email=email, role="Técnico", department="TI",
    ))
    MovementService.create_movement(db, MovementCreate(
        asset_id=asset.id,
        movement_type=MovementType.ALLOCATION,
        destination_custodian_id=cust.id,
        reason="Cautela (teste 034)",
        operator_name="Admin TI",
    ))
    db.refresh(asset)
    return cust


def test_new_inventory_snapshot_has_no_expected_custodian(db_session):
    """034/US1-1: novo inventário NÃO grava expected_custodian_name — nem p/ bem COM custodiante."""
    loc = _make_location(db_session, name="Loc 034 Resp", department="TI")
    asset = _make_asset(db_session, "INV-034-0001", location=loc)
    _assign_custodian(db_session, asset, "Maria Custodia", "maria.034@empresa.local")
    assert asset.custodian is not None  # pré-condição: bem com custodiante

    inv = InventarioService.create_inventario(db_session, name="Snapshot Sem Resp", location_id=loc.id)
    item = inv.itens[0]

    assert item.expected_custodian_name is None
    # Snapshot continua com tombamento + local esperado
    assert item.asset_id == asset.id
    assert item.expected_location_id == loc.id
    assert item.expected_location_name == loc.name


def test_asset_without_custodian_enters_snapshot_normally(db_session):
    """034/US1-2: bem sem custodiante entra normalmente no snapshot."""
    loc = _make_location(db_session, name="Loc 034 Livre", department="TI")
    asset = _make_asset(db_session, "INV-034-0002", location=loc)  # sem custodiante

    inv = InventarioService.create_inventario(db_session, name="Snapshot Sem Custodia", location_id=loc.id)
    item = inv.itens[0]

    assert item.asset_id == asset.id
    assert item.expected_custodian_name is None
    assert item.expected_location_name == loc.name
    assert item.status == InventarioItemStatus.PENDING


def test_snapshot_immutable_after_registration_changes(db_session):
    """034/US1-3: alterar custodiante/local do bem depois NÃO altera o snapshot."""
    loc_a = _make_location(db_session, name="Loc 034 A", department="TI")
    loc_b = _make_location(db_session, name="Loc 034 B", department="RH")
    asset = _make_asset(db_session, "INV-034-0003", location=loc_a)
    inv = InventarioService.create_inventario(db_session, name="Imutavel 034", location_id=loc_a.id)

    _assign_custodian(db_session, asset, "Novo Custodia", "novo.034@empresa.local")
    asset.location_id = loc_b.id
    db_session.commit()
    db_session.expire_all()

    item = InventarioService.get_by_id(db_session, inv.id).itens[0]
    assert item.expected_custodian_name is None
    assert item.expected_location_id == loc_a.id
    assert item.expected_location_name == loc_a.name


# ============================================================================
# Feature 034 — Conferência e ata sem "responsável esperado" (US2)
# ============================================================================

def _legacy_custodian_value(db, item, value):
    """Simula inventário legado: grava o snapshot textual direto no banco."""
    item.expected_custodian_name = value
    db.commit()


def test_conference_ignores_custodian_difference(db_session):
    """034/US2-1: responsável diferente não gera divergência (conformidade por local)."""
    loc = _make_location(db_session, name="Loc 034 Conf", department="TI")
    asset = _make_asset(db_session, "INV-034-0010", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Conf Resp", location_id=loc.id)
    item = inv.itens[0]
    _legacy_custodian_value(db_session, item, "Maria")

    # Cadastro do bem mudou de custodiante depois do snapshot (via fluxo de alocação)
    _assign_custodian(db_session, asset, "João", "joao.034@empresa.local")

    result = InventarioService.record_check(
        db_session, item=item, result=InventarioItemStatus.FOUND,
        user_id=1, username="conferente",
    )
    assert result.status == InventarioItemStatus.FOUND
    assert result.status != InventarioItemStatus.FOUND_WRONG_LOCATION
    assert result.status != InventarioItemStatus.NOT_FOUND


def test_location_divergence_still_detected(db_session):
    """034/US2-2: divergência de LOCAL continua detectada (FOUND_WRONG_LOCATION)."""
    loc = _make_location(db_session, name="Loc 034 Div", department="TI")
    _make_asset(db_session, "INV-034-0011", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Conf Local", location_id=loc.id)
    other_loc = _make_location(db_session, name="Loc 034 Outro", department="RH")

    InventarioService.record_check(
        db_session, item=inv.itens[0],
        result=InventarioItemStatus.FOUND_WRONG_LOCATION,
        found_location_id=other_loc.id,
    )
    assert inv.itens[0].status == InventarioItemStatus.FOUND_WRONG_LOCATION


def test_ata_csv_new_inventory_has_no_custodian_column(db_session):
    """034/US2-3 (FR-007): ata CSV de inventário novo NÃO tem a coluna "Responsável Esperado"."""
    loc = _make_location(db_session, name="Loc 034 CSV Novo", department="TI")
    _make_asset(db_session, "INV-034-0012", location=loc)
    inv = InventarioService.create_inventario(db_session, name="CSV Novo 034", location_id=loc.id)

    csv_text = ReportService.generate_inventario_csv(db_session, inv)
    assert "Local Esperado" in csv_text          # colunas restantes intactas
    assert "Responsável Esperado" not in csv_text


def test_ata_csv_legacy_inventory_keeps_custodian_column(db_session):
    """034/US2-4 (H-2/D2): ata de inventário legado preserva coluna e histórico."""
    loc = _make_location(db_session, name="Loc 034 CSV Leg", department="TI")
    _make_asset(db_session, "INV-034-0013", location=loc)
    _make_asset(db_session, "INV-034-0014", location=loc)
    inv = InventarioService.create_inventario(db_session, name="CSV Legado 034", location_id=loc.id)
    itens = {i.asset.tag: i for i in inv.itens}
    _legacy_custodian_value(db_session, itens["INV-034-0013"], "Maria Legado")
    # INV-034-0014 permanece sem valor → linha "Estoque / Livre" (contrato §1)

    csv_text = ReportService.generate_inventario_csv(db_session, inv)
    assert "Responsável Esperado" in csv_text
    assert "Maria Legado" in csv_text
    assert "Estoque / Livre" in csv_text


def test_ata_excel_new_no_column_legacy_with_column(db_session):
    """034/US2-5: Excel novo sem a coluna; legado com a coluna e histórico."""
    from openpyxl import load_workbook
    import io as _io

    # Novo
    loc = _make_location(db_session, name="Loc 034 XLSX", department="TI")
    _make_asset(db_session, "INV-034-0015", location=loc)
    inv_new = InventarioService.create_inventario(db_session, name="XLSX Novo 034", location_id=loc.id)
    wb = load_workbook(filename=_io.BytesIO(ReportService.generate_inventario_excel(db_session, inv_new)))
    values_new = {str(c.value) for row in wb.active.iter_rows() for c in row if c.value is not None}
    assert "Local Esperado" in values_new
    assert "Responsável Esperado" not in values_new

    # Legado
    _make_asset(db_session, "INV-034-0016", location=loc)
    inv_leg = InventarioService.create_inventario(db_session, name="XLSX Legado 034", location_id=loc.id)
    _legacy_custodian_value(db_session, inv_leg.itens[0], "Maria Legado XLSX")
    wb = load_workbook(filename=_io.BytesIO(ReportService.generate_inventario_excel(db_session, inv_leg)))
    values_leg = {str(c.value) for row in wb.active.iter_rows() for c in row if c.value is not None}
    assert "Responsável Esperado" in values_leg
    assert "Maria Legado XLSX" in values_leg


def _pdf_stream_text(pdf_bytes):
    """Extrai o conteúdo dos streams do PDF (ReportLab: ASCII85 + FlateDecode; stdlib)."""
    import re as _re
    import zlib as _zlib
    import base64 as _base64
    chunks = []
    for m in _re.finditer(rb"stream\r?\n(.*?)endstream", pdf_bytes, _re.DOTALL):
        data = m.group(1).strip()
        if data.endswith(b"~>"):  # camada ASCII85 do ReportLab
            try:
                data = _base64.a85decode(data, adobe=True)
            except Exception:
                pass
        try:
            data = _zlib.decompress(data)
        except Exception:
            pass
        chunks.append(data)
    return b"\n".join(chunks)


def test_ata_pdf_new_no_column_legacy_with_column(db_session):
    """034/US2-6: PDF novo sem a coluna; legado com a coluna e histórico."""
    # Novo
    loc = _make_location(db_session, name="Loc 034 PDF", department="TI")
    _make_asset(db_session, "INV-034-0017", location=loc)
    inv_new = InventarioService.create_inventario(db_session, name="PDF Novo 034", location_id=loc.id)
    pdf_new = ReportService.generate_inventario_pdf(db_session, inv_new)
    assert pdf_new.startswith(b"%PDF-")
    text_new = _pdf_stream_text(pdf_new)
    assert b"Respons" not in text_new  # sem "Responsável Esperado"

    # Legado
    _make_asset(db_session, "INV-034-0018", location=loc)
    inv_leg = InventarioService.create_inventario(db_session, name="PDF Legado 034", location_id=loc.id)
    _legacy_custodian_value(db_session, inv_leg.itens[0], "Maria Legado PDF")
    text_leg = _pdf_stream_text(ReportService.generate_inventario_pdf(db_session, inv_leg))
    assert b"Respons" in text_leg       # coluna presente no legado
    assert b"Maria Legado PDF" in text_leg


def test_templates_do_not_render_expected_custodian(client, db_session):
    """034/US2-7 (D3): telas do inventário não exibem colaborador esperado — nem em legado."""
    loc = _make_location(db_session, name="Loc 034 Tela", department="TI")
    asset = _make_asset(db_session, "INV-034-0019", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Telas 034", location_id=loc.id)
    item = inv.itens[0]
    _legacy_custodian_value(db_session, item, "Maria Tela")

    detail = client.get(f"/inventarios/{inv.id}")
    assert detail.status_code == 200
    assert "Maria Tela" not in detail.text
    assert "Colaborador esperado" not in detail.text

    conferir = client.get(f"/inventarios/{inv.id}/conferir/{asset.id}")
    assert conferir.status_code == 200
    assert "Maria Tela" not in conferir.text
    assert "Colaborador esperado" not in conferir.text


# ============================================================================
# Feature 034 — Inventários históricos íntegros (US3)
# ============================================================================

def test_legacy_inventory_consultable_with_history(db_session):
    """034/US3-1 (H-1): inventário legado permanece consultável; ata CSV preserva histórico."""
    loc = _make_location(db_session, name="Loc 034 Legado", department="TI")
    _make_asset(db_session, "INV-034-0020", location=loc)
    _make_asset(db_session, "INV-034-0021", location=loc)
    inv = InventarioService.create_inventario(db_session, name="Legado 034", location_id=loc.id)
    itens = {i.asset.tag: i for i in inv.itens}
    _legacy_custodian_value(db_session, itens["INV-034-0020"], "Maria Historico")

    for item in itens.values():
        InventarioService.record_check(
            db_session, item=item, result=InventarioItemStatus.FOUND,
            user_id=1, username="conferente",
        )

    s = InventarioService.summary(db_session, inv.id)
    assert s["expected"] == 2
    assert s["found"] == 2
    assert s["checked"] == 2

    csv_text = ReportService.generate_inventario_csv(db_session, inv)
    assert "Responsável Esperado" in csv_text
    assert "Maria Historico" in csv_text
